"""Uçtan uca duman testi (oku → sınıflandır → Excel) ve pencerenin açılıp kapanması.  python test_akis.py"""
import os
import sys
import tempfile
import textwrap

from openpyxl import load_workbook

from mailsinif import oku, rapor
from mailsinif.motor import Motor

sys.stdout.reconfigure(encoding="utf-8")

with tempfile.TemporaryDirectory() as d:
    # .eml (düz metin), .eml (yalnız HTML, Türkçe karakter), .mbox, bozuk dosya
    open(os.path.join(d, "a.eml"), "w", encoding="utf-8").write(textwrap.dedent("""\
        From: musteri@ornek.test
        Subject: =?utf-8?q?Fatura_hakk=C4=B1nda?=
        Content-Type: text/plain; charset=utf-8

        Merhaba, fatura ve dekont hakkında bilgi almak istiyorum, ödeme yaptım.
        """))
    open(os.path.join(d, "b.eml"), "w", encoding="utf-8").write(textwrap.dedent("""\
        From: x@ornek.test
        Subject: Kampanya
        Content-Type: text/html; charset=utf-8

        <html><body><p>SON 24 SAAT <b>indirim</b>!! <a href="x">tıklayın</a></p></body></html>
        """))
    open(os.path.join(d, "c.mbox"), "w", encoding="utf-8").write(
        "From a@b Mon Jan 1 00:00:00 2026\nFrom: a@b.test\nSubject: Siparis\n\nSiparişimi iptal etmek istiyorum.\n\n"
        "From c@d Mon Jan 1 00:00:01 2026\nFrom: c@d.test\nSubject: Soru\n\nStokta var mı, kaç gün sürer?\n")
    open(os.path.join(d, "bozuk.csv"), "wb").write(b"\xff\xfe\x00bozuk")
    os.makedirs(os.path.join(d, "alt"))
    with open(os.path.join("ornek_mailler", "yeni_mailler.csv"), encoding="utf-8") as f, \
            open(os.path.join(d, "alt", "yeni.csv"), "w", encoding="utf-8") as g:
        g.write(f.read())

    mailler, hatalar = oku.oku_klasor(d)
    assert len(mailler) == 2 + 2 + 16, len(mailler)
    assert any("Fatura hakkında" == m["konu"] for m in mailler), "başlık çözümlenmedi"
    html_mail = next(m for m in mailler if m["konu"] == "Kampanya")
    assert "<" not in html_mail["govde"] and "indirim" in html_mail["govde"], html_mail["govde"]
    print("okuma tamam:", len(mailler), "mail,", len(hatalar), "okunamayan:", hatalar)

    motor = Motor.egit(kaydet=False)
    sonuc = [{**m, **motor.siniflandir(m["konu"], m["govde"])} for m in mailler]
    yol = os.path.join(d, "rapor.xlsx")
    rapor.yaz(sonuc, yol, hatalar)
    wb = load_workbook(yol)
    assert wb.sheetnames[:3] == ["Özet", "Mailler", "Kontrol listesi"], wb.sheetnames
    assert wb["Mailler"].max_row == len(mailler) + 1
    print("Excel tamam:", wb.sheetnames)

    # elle yazılmış 16 yeni mail: beklenen kategoriler (eğitim setinde yok)
    beklenen = ["Sipariş / Talep", "Fatura / Ödeme", "Şikayet / Sorun", "Soru / Bilgi", "Belge / Evrak",
                "Otomatik bildirim", "Reklam / Spam", "Belirsiz", "Fatura / Ödeme", "Şikayet / Sorun",
                "Sipariş / Talep", "Reklam / Spam", "Soru / Bilgi", "Otomatik bildirim", "Şikayet / Sorun",
                "Otomatik bildirim"]
    yeni = [s for s in sonuc if s["kaynak"].startswith("yeni.csv")]
    dogru = sum(s["kategori"] == b for s, b in zip(yeni, beklenen))
    print(f"16 yeni elle yazılmış mail: {dogru}/16 doğru")
    for s, b in zip(yeni, beklenen):
        if s["kategori"] != b:
            print(f"  yanlış: {s['konu']!r} beklenen={b} tahmin={s['kategori']} güven={s['guven']}")

# pencere açılıp kapanıyor mu
import tkinter as tk
from uygulama import Uygulama

try:
    import time
    app = Uygulama()
    app.update()
    app.klasor.set("ornek_mailler")
    app.baslat()  # "Sınıflandır" düğmesinin işi: arka plan iş parçacığı + kuyruk
    bitis = time.time() + 120
    while app.calisiyor and time.time() < bitis:
        app.update()
        time.sleep(0.05)
    assert not app.calisiyor and len(app.sonuclar) == 16, (app.calisiyor, len(app.sonuclar))
    assert len(app.tablo.get_children()) == 16 and str(app.excel_dugme["state"]) == "normal"
    print("pencere tamam:", app.ozet.get())
    app.destroy()
except tk.TclError as e:
    print("pencere test edilemedi (ekran yok):", e)
