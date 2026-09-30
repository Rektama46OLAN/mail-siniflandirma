"""Excel raporu: Özet, Mailler ve Kontrol listesi (müşterinin doğru kategoriyi seçebileceği açılır liste ile)."""
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .kural import KATEGORILER

_BASLIK = PatternFill("solid", fgColor="1F3864")
_ACIL = PatternFill("solid", fgColor="FCE4D6")
_KONTROL = PatternFill("solid", fgColor="FFF2CC")


def _baslik(ws, kolonlar, genislikler):
    ws.append(kolonlar)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _BASLIK
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, g in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A2"


def yaz(sonuclar: list[dict], yol: str, okunamayan: list[str] | None = None):
    """sonuclar: her biri kaynak, gonderen, konu, govde + kategori, guven, acil, kontrol alanlarına sahip."""
    n = len(sonuclar)
    sayac = Counter(s["kategori"] for s in sonuclar)
    wb = Workbook()

    ws = wb.active
    ws.title = "Özet"
    ws["A1"], ws["A1"].font = "Mail Sınıflandırma Raporu", Font(bold=True, size=14)
    ws.append([])
    ws.append(["Toplam mail", n])
    ws.append(["Sınıflandırılan", n - sayac.get("Belirsiz", 0)])
    ws.append(["Belirsiz (insana bırakıldı)", sayac.get("Belirsiz", 0)])
    ws.append(["Acil işaretli", sum(s["acil"] for s in sonuclar)])
    ws.append(["Kontrol listesinde", sum(s["kontrol"] for s in sonuclar)])
    if okunamayan:
        ws.append(["Okunamayan dosya", len(okunamayan)])
    ws.append([])
    ws.append(["Kategori", "Adet", "Oran"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for k in KATEGORILER:
        ws.append([k, sayac.get(k, 0), (sayac.get(k, 0) / n) if n else 0])
        ws.cell(ws.max_row, 3).number_format = "0.0%"
    ws.append([])
    ws.append(["Not: kategori tahminleri otomatiktir; 'Kontrol listesi' sayfasındaki mailleri gözden geçirin."])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12

    ws = wb.create_sheet("Mailler")
    _baslik(ws, ["Dosya", "Gönderen", "Konu", "Kategori", "Güven", "Acil", "Kontrol", "Önizleme"],
            [24, 28, 38, 20, 8, 7, 9, 70])
    for s in sonuclar:
        ws.append([s["kaynak"], s["gonderen"], s["konu"], s["kategori"], s["guven"], "ACİL" if s["acil"] else "",
                   "kontrol et" if s["kontrol"] else "", " ".join(s["govde"].split())[:200]])
        satir = ws.max_row
        ws.cell(satir, 5).number_format = "0%"
        if s["acil"]:
            for c in ws[satir]:
                c.fill = _ACIL
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Kontrol listesi")
    _baslik(ws, ["Dosya", "Konu", "Tahmin", "Güven", "Acil", "Doğru kategori (siz seçin)", "Önizleme"],
            [24, 38, 20, 8, 7, 28, 70])
    dv = DataValidation(type="list", formula1='"' + ",".join(k for k in KATEGORILER) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for s in sorted((s for s in sonuclar if s["kontrol"]), key=lambda s: s["guven"]):
        ws.append([s["kaynak"], s["konu"], s["kategori"], s["guven"], "ACİL" if s["acil"] else "", None,
                   " ".join(s["govde"].split())[:200]])
        satir = ws.max_row
        ws.cell(satir, 4).number_format = "0%"
        ws.cell(satir, 6).fill = _KONTROL
        dv.add(ws.cell(satir, 6))

    if okunamayan:
        ws = wb.create_sheet("Okunamayanlar")
        ws.append(["Okunamayan dosyalar"])
        for h in okunamayan:
            ws.append([h])
        ws.column_dimensions["A"].width = 100
    wb.save(yol)
